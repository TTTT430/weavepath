from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from api.app import create_app
from api.llm import DisabledLLM
from graph_core import Conflict, GraphStore
from graph_core.attachments import parse_attachment


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Route memory", level=1)
    document.add_paragraph("Only this conversation may use the private requirement.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cases"
    sheet.append(["id", "expected"])
    sheet.append(["case-1", "positive"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box.text_frame.text = "Experiment acceptance criterion"
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def test_office_parsers_preserve_human_source_locators(tmp_path: Path):
    samples = [
        ("notes.docx", _docx_bytes(), "python-docx", "document", "private requirement"),
        ("cases.xlsx", _xlsx_bytes(), "openpyxl", "sheet Cases", "case-1"),
        ("brief.pptx", _pptx_bytes(), "python-pptx", "slide 1", "acceptance criterion"),
    ]
    for name, content, expected_parser, expected_locator, expected_text in samples:
        path = tmp_path / name
        path.write_bytes(content)
        parser, chunks = parse_attachment(path, name, "application/octet-stream")
        assert parser == expected_parser
        assert expected_locator in chunks[0]["locator"]
        assert expected_text in " ".join(chunk["content"] for chunk in chunks)


def test_disk_assets_route_listing_and_bound_citations_are_isolated(tmp_path: Path):
    database = tmp_path / "workspace.db"
    store = GraphStore(database)
    app = create_app(store, DisabledLLM())
    with TestClient(app) as client:
        graph = client.post("/api/v1/workflows", json={
            "name": "Files", "rootTitle": "A", "rootInstanceId": "A",
        }).json()
        workflow_id = graph["workflowId"]
        client.post(
            f"/api/v1/workflows/{workflow_id}/instances/A/fork",
            json={"title": "B", "instanceId": "B"},
        )
        client.post(
            f"/api/v1/workflows/{workflow_id}/instances/A/fork",
            json={"title": "E", "instanceId": "E"},
        )
        uploaded = client.post(
            f"/api/v1/workflows/{workflow_id}/instances/B/attachments"
            "?name=private.docx&mimeType=application%2Fvnd.openxmlformats-officedocument.wordprocessingml.document",
            content=_docx_bytes(),
        )
        assert uploaded.status_code == 201
        attachment = uploaded.json()
        assert attachment["parseStatus"] == "processing"
        attachment = client.get(
            f"/api/v1/workflows/{workflow_id}/instances/B/attachments/"
            f"{attachment['attachmentId']}"
        ).json()
        assert attachment["parseStatus"] == "ready"
        assert attachment["parser"] == "python-docx"
        assert attachment["chunkCount"] >= 1

        object_path = database.parent / "files" / "objects" / attachment["sha256"][:2] / attachment["sha256"]
        assert object_path.is_file()
        row = store._conn.execute(
            "SELECT content_text,storage_key FROM message_attachments WHERE id=?",
            (attachment["attachmentId"],),
        ).fetchone()
        assert row["content_text"] == ""
        assert row["storage_key"] == attachment["sha256"]

        own_search_payload = client.get(
            f"/api/v1/workflows/{workflow_id}/instances/B/attachments/search",
            params={"q": "private requirement"},
        ).json()
        assert own_search_payload["engine"] == "fts5-trigram"
        own_search = own_search_payload["results"]
        assert own_search[0]["attachmentId"] == attachment["attachmentId"]
        assert own_search[0]["routeInstanceId"] == "B"
        assert own_search[0]["inherited"] is False
        assert "private requirement" in own_search[0]["preview"]
        assert own_search[0]["chunkSha256"]
        index_status = client.get(
            f"/api/v1/workflows/{workflow_id}/instances/B/attachments/search-status"
        ).json()
        assert index_status["ready"] is True
        assert index_status["engine"] == "fts5-trigram"
        assert index_status["indexedChunks"] == index_status["readyChunks"] >= 1
        sibling_search = client.get(
            f"/api/v1/workflows/{workflow_id}/instances/E/attachments/search",
            params={"q": "private requirement"},
        ).json()["results"]
        assert sibling_search == []

        envelope = "[WeavePath attachments v2]\n" + json.dumps({
            "files": [{
                "attachmentId": attachment["attachmentId"],
                "name": attachment["name"],
                "mimeType": attachment["mimeType"],
                "size": attachment["size"],
            }],
            "prompt": "Find the private requirement",
        })
        assert client.post(
            f"/api/v1/workflows/{workflow_id}/instances/B/messages",
            json={"role": "user", "content": envelope},
        ).status_code == 201

        local = client.get(
            f"/api/v1/workflows/{workflow_id}/instances/B/attachments?scope=local"
        ).json()["attachments"]
        assert local[0]["contextSources"][0]["locator"].startswith("document")
        assert local[0]["inherited"] is False
        assert client.get(
            f"/api/v1/workflows/{workflow_id}/instances/E/attachments?scope=route"
        ).json()["attachments"] == []

        client.post(
            f"/api/v1/workflows/{workflow_id}/instances/B/fork",
            json={"title": "C", "instanceId": "C"},
        )
        inherited = client.get(
            f"/api/v1/workflows/{workflow_id}/instances/C/attachments?scope=route"
        ).json()["attachments"]
        assert inherited[0]["attachmentId"] == attachment["attachmentId"]
        assert inherited[0]["inherited"] is True
        inherited_search = client.get(
            f"/api/v1/workflows/{workflow_id}/instances/C/attachments/search",
            params={"q": "private requirement"},
        ).json()["results"]
        assert inherited_search[0]["attachmentId"] == attachment["attachmentId"]
        assert inherited_search[0]["inherited"] is True
        detail = client.get(
            f"/api/v1/workflows/{workflow_id}/instances/C/attachments/{attachment['attachmentId']}"
        ).json()
        assert detail["chunks"][0]["preview"]

    store.close()


def test_persistent_trigram_index_supports_chinese_short_fallback_and_deletion(tmp_path: Path):
    database = tmp_path / "workspace.db"
    store = GraphStore(database)
    workflow = store.create_workflow(
        name="Search", root_title="A", root_instance_id="A"
    )
    attachment = store.create_attachment(
        workflow["workflowId"], "A", name="中文资料.txt", mime_type="text/plain",
        size_bytes=len("中文情感分析证据编号 WP-7788".encode("utf-8")),
        content_text="中文情感分析证据编号 WP-7788",
    )
    full = store.search_attachments(
        workflow["workflowId"], "A", query="情感分析"
    )
    assert full["engine"] == "fts5-trigram"
    assert full["results"][0]["attachmentId"] == attachment["attachmentId"]
    short = store.search_attachments(workflow["workflowId"], "A", query="中文")
    assert short["engine"] == "route-like-short-query"
    assert short["results"][0]["attachmentId"] == attachment["attachmentId"]
    store.close()

    reopened = GraphStore(database)
    assert reopened.attachment_search_status(
        workflow["workflowId"], "A"
    )["ready"] is True
    assert reopened.search_attachments(
        workflow["workflowId"], "A", query="WP-7788"
    )["results"][0]["attachmentId"] == attachment["attachmentId"]
    reopened.delete_attachment(workflow["workflowId"], "A", attachment["attachmentId"])
    assert reopened.search_attachments(
        workflow["workflowId"], "A", query="WP-7788"
    )["results"] == []
    assert reopened.attachment_search_status(
        workflow["workflowId"], "A"
    )["indexedChunks"] == 0
    reopened.close()


def test_automatic_retrieval_plan_is_route_scoped_frozen_and_restart_safe(tmp_path: Path):
    database = tmp_path / "workspace.db"
    store = GraphStore(database)
    graph = store.create_workflow(
        name="Retrieval", root_title="A", root_instance_id="A"
    )
    workflow_id = graph["workflowId"]
    ancestor = store.create_attachment(
        workflow_id, "A", name="shared-notes.txt", mime_type="text/plain",
        size_bytes=80,
        content_text="Shared sentiment evidence marker ROUTE-ALPHA belongs to A.",
    )
    store.fork(workflow_id, "A", title="B", instance_id="B")
    store.fork(workflow_id, "B", title="C", instance_id="C")
    store.fork(workflow_id, "A", title="E", instance_id="E")
    sibling = store.create_attachment(
        workflow_id, "E", name="sibling-secret.txt", mime_type="text/plain",
        size_bytes=80,
        content_text="Sibling sentiment evidence marker ROUTE-EPSILON is private.",
    )

    user = store.append_message(
        workflow_id, "C", role="user",
        content="Compare the sentiment evidence marker for this route.",
    )
    plan = store.retrieval_plan_for_message(workflow_id, user["id"])
    assert plan is not None
    assert plan["mode"] == "automatic"
    assert plan["engine"] == "fts5-trigram"
    assert plan["routeInstanceIds"] == ["A", "B", "C"]
    assert plan["selectedChunks"] == 1
    assert len(plan["contextText"]) <= plan["budgetCharacters"]
    assert plan["sources"][0]["attachmentId"] == ancestor["attachmentId"]
    assert plan["sources"][0]["inherited"] is True
    assert sibling["attachmentId"] not in str(plan)
    materialized = store.materialize_messages(
        workflow_id, store.list_messages(workflow_id, "C")["messages"]
    )
    assert "ROUTE-ALPHA" in materialized[-1]["content"]
    assert "ROUTE-EPSILON" not in materialized[-1]["content"]
    frozen_hash = plan["contextSha256"]
    with pytest.raises(Conflict, match="retrieval plan"):
        store.delete_attachment(workflow_id, "A", ancestor["attachmentId"])
    with pytest.raises(Conflict, match="retrieval plan"):
        store.reparse_attachment(workflow_id, "A", ancestor["attachmentId"])
    store.close()

    reopened = GraphStore(database)
    restored = reopened.retrieval_plan_for_message(workflow_id, user["id"])
    assert restored is not None
    assert restored["contextSha256"] == frozen_hash
    assert restored["sources"][0]["chunkSha256"] == plan["sources"][0]["chunkSha256"]
    materialized = reopened.materialize_messages(
        workflow_id, reopened.list_messages(workflow_id, "C")["messages"]
    )
    assert "ROUTE-ALPHA" in materialized[-1]["content"]
    assert "ROUTE-EPSILON" not in materialized[-1]["content"]
    reopened.close()


def test_automatic_retrieval_plan_enforces_chunk_and_character_budgets():
    store = GraphStore(":memory:")
    graph = store.create_workflow(
        name="Budget", root_title="A", root_instance_id="A"
    )
    workflow_id = graph["workflowId"]
    for index in range(8):
        content = f"budget evidence marker file {index} " + ("x" * 5_000)
        store.create_attachment(
            workflow_id, "A", name=f"budget-{index}.txt", mime_type="text/plain",
            size_bytes=len(content.encode()), content_text=content,
        )

    plan = store.automatic_retrieval_plan(
        workflow_id, "A", "Review the budget evidence marker."
    )
    assert plan is not None
    assert 1 <= plan["selectedChunks"] <= 6
    assert len(plan["sources"]) == plan["selectedChunks"]
    assert len(plan["contextText"]) <= plan["budgetCharacters"] == 24_000
    assert plan["selectedCharacters"] <= plan["budgetCharacters"]
    assert plan["truncated"] is True
    store.close()


def test_image_upload_is_retained_with_an_explicit_ocr_status(tmp_path: Path):
    store = GraphStore(":memory:", attachment_root=tmp_path / "files")
    with TestClient(create_app(store, DisabledLLM())) as client:
        graph = client.post("/api/v1/workflows", json={
            "name": "Images", "rootTitle": "A", "rootInstanceId": "A",
        }).json()
        response = client.post(
            f"/api/v1/workflows/{graph['workflowId']}/instances/A/attachments"
            "?name=scan.png&mimeType=image%2Fpng",
            content=b"not-a-real-image-but-still-an-untrusted-asset",
        )
        assert response.status_code == 201
        assert response.json()["parseStatus"] == "processing"
        detail = client.get(
            f"/api/v1/workflows/{graph['workflowId']}/instances/A/attachments/"
            f"{response.json()['attachmentId']}"
        ).json()
        assert detail["parseStatus"] == "failed"
        assert detail["parseErrorCode"] == "attachmentOcrUnavailable"
    store.close()


def test_processing_asset_is_resumed_after_a_service_restart(tmp_path: Path):
    database = tmp_path / "workspace.db"
    store = GraphStore(database)
    workflow = store.create_workflow(
        name="Recovery", root_title="A", root_instance_id="A"
    )
    staged = store.new_attachment_staging_path()
    content = b"durable parser recovery"
    staged.write_bytes(content)
    uploaded = store.create_attachment_from_file(
        workflow["workflowId"], "A", name="recovery.txt", mime_type="text/plain",
        size_bytes=len(content), staged_path=staged,
        sha256=hashlib.sha256(content).hexdigest(), parse_immediately=False,
    )
    assert uploaded["parseStatus"] == "processing"
    store.close()

    reopened = GraphStore(database)
    detail = reopened.get_attachment(
        workflow["workflowId"], "A", uploaded["attachmentId"]
    )
    assert detail["parseStatus"] == "ready"
    assert detail["chunks"][0]["preview"] == "durable parser recovery"
    reopened.close()


def test_startup_sweep_removes_only_unreferenced_objects_and_staging_files(tmp_path: Path):
    database = tmp_path / "workspace.db"
    files = tmp_path / "files"
    orphan_hash = "a" * 64
    orphan = files / "objects" / orphan_hash[:2] / orphan_hash
    orphan.parent.mkdir(parents=True)
    orphan.write_bytes(b"orphan")
    staging = files / "staging" / "abandoned.upload"
    staging.parent.mkdir(parents=True)
    staging.write_bytes(b"partial")

    store = GraphStore(database, attachment_root=files)
    assert not orphan.exists()
    assert not staging.exists()
    store.close()


def test_chunk_upload_resumes_after_restart_and_assembles_exact_bytes(tmp_path: Path):
    database = tmp_path / "workspace.db"
    chunk_size = 256 * 1024
    content = b"a" * chunk_size + b"tail-of-file"
    store = GraphStore(database)
    with TestClient(create_app(store, DisabledLLM())) as client:
        workflow = client.post("/api/v1/workflows", json={
            "name": "Resumable", "rootTitle": "A", "rootInstanceId": "A",
        }).json()
        route = f"/api/v1/workflows/{workflow['workflowId']}/instances/A"
        session = client.post(route + "/attachment-uploads", json={
            "clientKey": "browser-file-fingerprint", "name": "large.txt",
            "mimeType": "text/plain", "size": len(content), "chunkSize": chunk_size,
        })
        assert session.status_code == 201
        upload = session.json()
        assert upload["missingChunks"] == [0, 1]
        second = client.put(
            route + f"/attachment-uploads/{upload['uploadId']}/chunks/1",
            content=content[chunk_size:],
        )
        assert second.status_code == 200
        assert second.json()["receivedChunks"] == [1]
        assert second.json()["missingChunks"] == [0]
        duplicate = client.put(
            route + f"/attachment-uploads/{upload['uploadId']}/chunks/1",
            content=content[chunk_size:],
        )
        assert duplicate.status_code == 200
    store.close()

    reopened = GraphStore(database)
    with TestClient(create_app(reopened, DisabledLLM())) as client:
        route = f"/api/v1/workflows/{workflow['workflowId']}/instances/A"
        resumed = client.post(route + "/attachment-uploads", json={
            "clientKey": "browser-file-fingerprint", "name": "large.txt",
            "mimeType": "text/plain", "size": len(content), "chunkSize": chunk_size,
        }).json()
        assert resumed["uploadId"] == upload["uploadId"]
        assert resumed["missingChunks"] == [0]
        assert client.put(
            route + f"/attachment-uploads/{upload['uploadId']}/chunks/0",
            content=content[:chunk_size],
        ).status_code == 200
        completed = client.post(
            route + f"/attachment-uploads/{upload['uploadId']}/complete"
        )
        assert completed.status_code == 200
        attachment = completed.json()
        assert attachment["parseStatus"] == "processing"
        detail = client.get(
            route + f"/attachments/{attachment['attachmentId']}"
        ).json()
        assert detail["parseStatus"] == "ready"
        object_path = (database.parent / "files" / "objects" /
                       attachment["sha256"][:2] / attachment["sha256"])
        assert object_path.read_bytes() == content
    reopened.close()
