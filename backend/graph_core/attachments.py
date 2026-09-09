from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable


MAX_ATTACHMENT_BYTES = 50 * 1024 * 1024
CHUNK_CHARACTERS = 6_000

TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "json", "jsonl", "csv", "tsv", "yaml", "yml",
    "xml", "html", "css", "js", "jsx", "ts", "tsx", "py", "java", "c", "h",
    "cpp", "hpp", "cs", "go", "rs", "rb", "php", "sh", "ps1", "sql", "toml",
    "ini", "cfg", "log", "tex", "r",
}
TEXT_MIME_TYPES = {
    "application/json", "application/ld+json", "application/xml", "application/yaml",
    "application/javascript", "application/x-javascript", "application/sql",
}
DOCUMENT_EXTENSIONS = {"pdf", "docx", "xlsx", "pptx"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "tif", "tiff", "bmp"}


class AttachmentParseError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def attachment_kind(name: str, mime_type: str) -> str | None:
    suffix = Path(name).suffix.lower().removeprefix(".")
    normalized_mime = mime_type.lower().split(";", 1)[0].strip()
    if (normalized_mime.startswith("text/") or normalized_mime in TEXT_MIME_TYPES
            or suffix in TEXT_EXTENSIONS):
        return "text"
    if suffix in DOCUMENT_EXTENSIONS:
        return suffix
    if normalized_mime == "application/pdf":
        return "pdf"
    if normalized_mime.startswith("image/") or suffix in IMAGE_EXTENSIONS:
        return "image"
    return None


def supports_attachment(name: str, mime_type: str) -> bool:
    return attachment_kind(name, mime_type) is not None


def _chunk_lines(lines: Iterable[str], locator_prefix: str) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    buffered: list[str] = []
    buffered_chars = 0
    start_line = 1
    line_number = 0

    def flush(end_line: int) -> None:
        nonlocal buffered, buffered_chars, start_line
        text = "".join(buffered).strip()
        if text:
            locator = f"{locator_prefix} · lines {start_line}-{end_line}"
            chunks.append({"locator": locator, "content": text})
        buffered, buffered_chars = [], 0
        start_line = end_line + 1

    for line_number, line in enumerate(lines, 1):
        value = line if line.endswith("\n") else line + "\n"
        if len(value) > CHUNK_CHARACTERS:
            if buffered:
                flush(line_number - 1)
            for offset in range(0, len(value), CHUNK_CHARACTERS):
                part = value[offset:offset + CHUNK_CHARACTERS].strip()
                if part:
                    chunks.append({
                        "locator": f"{locator_prefix} · line {line_number} · part "
                                   f"{offset // CHUNK_CHARACTERS + 1}",
                        "content": part,
                    })
            start_line = line_number + 1
            continue
        if buffered and buffered_chars + len(value) > CHUNK_CHARACTERS:
            flush(line_number - 1)
        if not buffered:
            start_line = line_number
        buffered.append(value)
        buffered_chars += len(value)
    if buffered:
        flush(line_number)
    return chunks


def _chunk_block(text: str, locator: str) -> list[dict[str, Any]]:
    if not text.strip():
        return []
    chunks = _chunk_lines(text.splitlines(keepends=True), locator)
    if len(chunks) == 1:
        chunks[0]["locator"] = locator
    return chunks


def _parse_text(path: Path) -> tuple[str, list[dict[str, Any]]]:
    try:
        text = path.read_bytes().decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AttachmentParseError(
            "attachmentUnreadable", "The file is not valid UTF-8 text."
        ) from exc
    if not text.strip() or "\x00" in text:
        raise AttachmentParseError(
            "attachmentUnreadable", "The file does not contain readable text."
        )
    return "utf8-text", _chunk_lines(text.splitlines(keepends=True), "text")


def _parse_pdf(path: Path) -> tuple[str, list[dict[str, Any]]]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        chunks: list[dict[str, Any]] = []
        for index, page in enumerate(reader.pages, 1):
            chunks.extend(_chunk_block(page.extract_text() or "", f"page {index}"))
    except Exception as exc:
        raise AttachmentParseError("attachmentParseFailed", "The PDF could not be parsed.") from exc
    if not chunks:
        raise AttachmentParseError(
            "attachmentOcrRequired",
            "No selectable text was found in the PDF. OCR is not configured yet.",
        )
    return "pypdf", chunks


def _parse_docx(path: Path) -> tuple[str, list[dict[str, Any]]]:
    from docx import Document

    try:
        document = Document(str(path))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(f"[Table {table_index}, row {row_index}] " + "\t".join(cells))
    except Exception as exc:
        raise AttachmentParseError("attachmentParseFailed", "The Word document could not be parsed.") from exc
    chunks = _chunk_lines((block + "\n" for block in blocks), "document")
    if not chunks:
        raise AttachmentParseError("attachmentUnreadable", "The Word document contains no readable text.")
    return "python-docx", chunks


def _parse_xlsx(path: Path) -> tuple[str, list[dict[str, Any]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        chunks: list[dict[str, Any]] = []
        try:
            for sheet in workbook.worksheets:
                lines: list[str] = []
                row_numbers: list[int] = []
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = ["" if value is None else str(value) for value in row]
                    if not any(values):
                        continue
                    lines.append("\t".join(values) + "\n")
                    row_numbers.append(row_number)
                sheet_chunks = _chunk_lines(lines, f"sheet {sheet.title}")
                # The generic line locator is deterministic; expose that these
                # are non-empty worksheet rows rather than claiming Excel's
                # sparse absolute row numbers when gaps exist.
                if row_numbers:
                    for chunk in sheet_chunks:
                        chunk["locator"] = chunk["locator"].replace("lines", "data rows")
                chunks.extend(sheet_chunks)
        finally:
            workbook.close()
    except Exception as exc:
        raise AttachmentParseError("attachmentParseFailed", "The spreadsheet could not be parsed.") from exc
    if not chunks:
        raise AttachmentParseError("attachmentUnreadable", "The spreadsheet contains no readable cells.")
    return "openpyxl", chunks


def _parse_pptx(path: Path) -> tuple[str, list[dict[str, Any]]]:
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
        chunks: list[dict[str, Any]] = []
        for slide_index, slide in enumerate(presentation.slides, 1):
            blocks: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and str(shape.text).strip():
                    blocks.append(str(shape.text).strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        values = [cell.text.strip() for cell in row.cells]
                        if any(values):
                            blocks.append("\t".join(values))
            chunks.extend(_chunk_block("\n".join(blocks), f"slide {slide_index}"))
    except Exception as exc:
        raise AttachmentParseError("attachmentParseFailed", "The presentation could not be parsed.") from exc
    if not chunks:
        raise AttachmentParseError("attachmentUnreadable", "The presentation contains no readable text.")
    return "python-pptx", chunks


def parse_attachment(path: str | Path, name: str, mime_type: str) -> tuple[str, list[dict[str, Any]]]:
    source = Path(path)
    kind = attachment_kind(name, mime_type)
    if kind == "text":
        return _parse_text(source)
    if kind == "pdf":
        return _parse_pdf(source)
    if kind == "docx":
        return _parse_docx(source)
    if kind == "xlsx":
        return _parse_xlsx(source)
    if kind == "pptx":
        return _parse_pptx(source)
    if kind == "image":
        raise AttachmentParseError(
            "attachmentOcrUnavailable",
            "Image OCR is not configured in this local preview.",
        )
    raise AttachmentParseError("attachmentUnsupported", "This file type is not supported.")
